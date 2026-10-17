import openpyxl
from pathlib import Path

# 创建工作簿
excel_file = openpyxl.Workbook()
# 创建工作表
sheet_obj = excel_file.create_sheet('人员信息', 0)
# 定义列标题
field_name = ['序号', '姓名', '电子邮件']
# 定义一组人员信息数据
rows = [
    [1, '刘峰', 'lfeng@163.com'],
    [2, '李明', 'lming@163.com'],
    [3, '王丽', '王丽@163.com'],
    [4, '石文', 'shweng@sina.com'],
    [5, '景小美', 'jxiaomei@126.com'],
    [6, '吴用', 'wuyong@sohu.com']
]
# 在第一行写入列标题
for i in range(len(field_name)):
    sheet_obj.cell(1, i+1, field_name[i])

# 从第二行开始写入人员信息，每行写入一个人的信息，一次写入
for j in range(len(rows)):
    for k in range(0, 3):
        sheet_obj.cell(j+2, k+1, rows[j][k])
        print(rows[j][k])

# 保存到文件中
excel_file.save(Path(__file__).parent / '人员信息表.xlsx')