import openpyxl
import pymysql
import os
from pathlib import Path

# 从数据库某张表中取出所有记录的函数
def get_rows(host, db_name, table_name, user, passwd):
    # 建立一个数据库的连接
    conn = pymysql.connect(host=host, port=3306, database=db_name, user=user, password=passwd)
    # 建立一个游标
    cur = conn.cursor()
    # 组合一个SQL查询语句
    sql = 'select * from ' + table_name
    # 执行SQL语句
    cur.execute(sql)
    # rows取得所有行记录，cur.fetchall()返回所有符合条件的记录
    rows = cur.fetchall()
    # cur.description返回数据表的字段信息
    # 返回值fields是一个元组，其中的每一项元素也是一个元组(子元组)
    # 这个子元组的第一个元素是字段名
    fields = cur.description
    # 关闭游标
    cur.close()
    return fields, rows


# 将表的记录导入Excel表的函数
def export_to_excel(host, db_name, table_name, user, passwd, filename):
    # 调用函数，取得数据表的字段信息和记录信息
    fields, table_rows = get_rows(host, db_name, table_name, user, passwd)
    # print(fields)
    # print(table_rows)
    # 生成Excel文件的工作簿
    workbook = openpyxl.Workbook()
    # 在工作簿中生成一个工作表，表名设为table_加数据表名
    sheet = workbook.create_sheet('table_' + table_name, 0)
    # 在工作表的第一行写上字段名
    for i in range(0, len(fields)):
        sheet.cell(1, i + 1, fields[i][0])
    # 从工作表第二行开始写入每条记录的内容
    for row in range(0, len(table_rows)):
        for col in range(0, len(fields)):
            sheet.cell(row+2, col+1, '%s' % table_rows[row][col])

    # 保存到Excel文件中
    workbook.save(filename)



# 主函数min
if __name__ == '__main__':
    # 初始化各变量值
    host = 'localhost'
    db_name = 'world'
    table_name = 'country'
    user = 'root'
    password = os.environ.get('DBPASSWORD')
    file_path = Path('__file__').parent / 'chapter06/coutries.xlsx';
    # 调用函数，将一张数据表的全部记录导入到一个Excel文件中
    export_to_excel(host, db_name, table_name, user, password, file_path)